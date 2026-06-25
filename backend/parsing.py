"""Parse trial balance input (file upload or pasted text) into normalized records.

Handles two real-world shapes in addition to generic CSV/pasted data:

1. Audit workpaper ("Working Trial Balance"): Code / Account / Description columns
   plus several balance columns (Unadjusted, Adjusted, Report <date>, ...). Category
   header rows and "<code> Total" subtotal rows must be excluded. Some exports include
   a hidden tag column (e.g. "System Type") whose values start with "Account_" for real
   account rows, "SubTotal_"/"Header_"/"Code_..." for structural rows, etc. — when present
   this is used to identify account rows precisely; otherwise we fall back to heuristics.

2. Client TB export: Full name / Debit / Credit columns, terminated by a "TOTAL" row.
   The signed balance is Debit - Credit.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass

import openpyxl
import pandas as pd
import pdfplumber


@dataclass
class TBEntry:
    account_name: str
    balance: float
    account_code: str | None = None


_AMOUNT_RE = re.compile(r"-?\(?\$?\s*-?[\d,]+\.?\d*\)?")


def _to_float(value) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text:
        return None
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()").replace("$", "").replace(",", "").strip()
    if not text:
        return None
    try:
        amount = float(text)
    except ValueError:
        return None
    return -amount if negative else amount


# ---------------------------------------------------------------------------
# Audit workpaper format
# ---------------------------------------------------------------------------

_TAG_ACCOUNT_PREFIXES = ("Account_",)
_TAG_NON_ACCOUNT_PREFIXES = (
    "Header_",
    "SubcodeBlankRow_",
    "SubTotal",
    "Code_",
    "GrandTotalRow",
    "NetIncomeRow",
    "TotalAssets",
    "TotalLiabilities",
    "TotalEquity",
    "TotalRevenue",
    "TotalExpense",
)


def _find_header_row(ws) -> int | None:
    for row in ws.iter_rows(min_row=1, max_row=min(30, ws.max_row)):
        values = [str(c.value).strip().lower() if c.value is not None else "" for c in row]
        if "code" in values and "account" in values and "description" in values:
            return row[0].row
    return None


def inspect_audit_workpaper(raw: bytes) -> dict:
    """Return detected balance columns so the caller can choose which to compare."""
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    ws = wb.worksheets[0]
    header_row_idx = _find_header_row(ws)
    if header_row_idx is None:
        return {"is_audit_workpaper": False, "columns": []}

    header_cells = next(ws.iter_rows(min_row=header_row_idx, max_row=header_row_idx))
    headers = [str(c.value).strip() if c.value is not None else "" for c in header_cells]

    balance_keywords = ("report", "adjusted", "unadjusted", "balance")
    candidates = [
        h
        for h in headers
        if h
        and any(k in h.lower() for k in balance_keywords)
        and "tie-out" not in h.lower()
        and "tie out" not in h.lower()
    ]

    # Prefer the right-most column whose header looks like "Report <date>".
    report_cols = [h for h in candidates if h.lower().startswith("report")]
    default = report_cols[-1] if report_cols else (candidates[-1] if candidates else None)

    return {"is_audit_workpaper": True, "columns": candidates, "default_column": default}


def _classify_rows_by_tag(ws, header_row_idx: int) -> list[int] | None:
    """If a tag column (e.g. 'System Type') exists, return row indices that are real accounts."""
    header_cells = next(ws.iter_rows(min_row=header_row_idx, max_row=header_row_idx))
    tag_col_idx = None
    for c in header_cells:
        if c.value and "system type" in str(c.value).strip().lower():
            tag_col_idx = c.column
            break
    if tag_col_idx is None:
        return None

    account_rows = []
    for row in ws.iter_rows(min_row=header_row_idx + 1, max_row=ws.max_row):
        tag_cell = row[tag_col_idx - 1]
        tag = tag_cell.value
        if isinstance(tag, str) and tag.startswith(_TAG_ACCOUNT_PREFIXES):
            account_rows.append(tag_cell.row)
    return account_rows


def parse_audit_workpaper(raw: bytes, balance_column: str | None = None) -> list[TBEntry]:
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    ws = wb.worksheets[0]
    header_row_idx = _find_header_row(ws)
    if header_row_idx is None:
        raise ValueError("Could not find a Code/Account/Description header row.")

    header_cells = next(ws.iter_rows(min_row=header_row_idx, max_row=header_row_idx))
    headers = {str(c.value).strip(): c.column for c in header_cells if c.value is not None}

    code_col = headers.get("Code")
    account_col = headers.get("Account")
    description_col = headers.get("Description")
    if description_col is None:
        raise ValueError("Audit workpaper is missing a Description column.")

    if balance_column is None:
        info = inspect_audit_workpaper(raw)
        balance_column = info.get("default_column")
    balance_col = headers.get(balance_column)
    if balance_col is None:
        raise ValueError(f"Balance column '{balance_column}' not found in workpaper.")

    account_rows = _classify_rows_by_tag(ws, header_row_idx)

    entries: list[TBEntry] = []
    if account_rows is not None:
        for row_idx in account_rows:
            description = ws.cell(row=row_idx, column=description_col).value
            balance = _to_float(ws.cell(row=row_idx, column=balance_col).value)
            code = ws.cell(row=row_idx, column=account_col).value if account_col else None
            if description and balance is not None:
                entries.append(
                    TBEntry(
                        account_name=str(description).strip(),
                        balance=balance,
                        account_code=str(code).strip() if code else None,
                    )
                )
    else:
        # Heuristic fallback: a real account row has a non-blank Account code and
        # a Description, and its Code column does NOT contain "Total".
        for row in ws.iter_rows(min_row=header_row_idx + 1, max_row=ws.max_row):
            code_val = row[code_col - 1].value if code_col else None
            account_val = row[account_col - 1].value if account_col else None
            description_val = row[description_col - 1].value
            if code_val and "total" in str(code_val).lower():
                continue
            if not account_val or not description_val:
                continue
            balance = _to_float(row[balance_col - 1].value)
            if balance is None:
                continue
            entries.append(
                TBEntry(
                    account_name=str(description_val).strip(),
                    balance=balance,
                    account_code=str(account_val).strip(),
                )
            )
    return entries


# ---------------------------------------------------------------------------
# Client TB (debit/credit) format
# ---------------------------------------------------------------------------


def _find_client_header_row(ws) -> int | None:
    for row in ws.iter_rows(min_row=1, max_row=min(20, ws.max_row)):
        values = [str(c.value).strip().lower() if c.value is not None else "" for c in row]
        has_name = any(v in ("full name", "account name", "account") for v in values)
        has_debit_credit = "debit" in values and "credit" in values
        if has_name and has_debit_credit:
            return row[0].row
    return None


def is_client_debit_credit_format(raw: bytes) -> bool:
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    ws = wb.worksheets[0]
    return _find_client_header_row(ws) is not None


def parse_client_debit_credit(raw: bytes) -> list[TBEntry]:
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    ws = wb.worksheets[0]
    header_row_idx = _find_client_header_row(ws)
    if header_row_idx is None:
        raise ValueError("Could not find a Full name/Debit/Credit header row.")

    header_cells = next(ws.iter_rows(min_row=header_row_idx, max_row=header_row_idx))
    headers = {str(c.value).strip().lower(): c.column for c in header_cells if c.value is not None}
    name_col = headers.get("full name") or headers.get("account name") or headers.get("account")
    debit_col = headers["debit"]
    credit_col = headers["credit"]

    entries: list[TBEntry] = []
    for row in ws.iter_rows(min_row=header_row_idx + 1, max_row=ws.max_row):
        name = row[name_col - 1].value
        if not name:
            continue
        name = str(name).strip()
        if name.upper() == "TOTAL":
            break
        debit = _to_float(row[debit_col - 1].value) or 0.0
        credit = _to_float(row[credit_col - 1].value) or 0.0
        entries.append(TBEntry(account_name=name, balance=debit - credit))
    return entries


# ---------------------------------------------------------------------------
# Generic fallback (arbitrary CSV/XLSX/PDF/pasted text with one balance column)
# ---------------------------------------------------------------------------


def _pick_columns(df: pd.DataFrame) -> tuple[str, str]:
    name_col = None
    balance_col = None
    for col in df.columns:
        label = str(col).strip().lower()
        if name_col is None and any(k in label for k in ("account", "name", "description")):
            name_col = col
        if balance_col is None and any(
            k in label for k in ("balance", "amount", "debit", "credit", "net")
        ):
            balance_col = col
    if name_col is None:
        name_col = df.columns[0]
    if balance_col is None:
        for col in reversed(df.columns):
            if col != name_col:
                balance_col = col
                break
    return name_col, balance_col


def _rows_to_entries(df: pd.DataFrame) -> list[TBEntry]:
    name_col, balance_col = _pick_columns(df)
    entries: list[TBEntry] = []
    for _, row in df.iterrows():
        name = str(row.get(name_col, "")).strip()
        if not name or name.lower() in ("nan", "total", "totals", "grand total"):
            continue
        balance = _to_float(row.get(balance_col))
        if balance is None:
            continue
        entries.append(TBEntry(account_name=name, balance=balance))
    return entries


def parse_csv(raw: bytes) -> list[TBEntry]:
    df = pd.read_csv(io.BytesIO(raw))
    return _rows_to_entries(df)


def parse_excel_generic(raw: bytes) -> list[TBEntry]:
    df = pd.read_excel(io.BytesIO(raw))
    return _rows_to_entries(df)


def parse_pdf(raw: bytes) -> list[TBEntry]:
    entries: list[TBEntry] = []
    with pdfplumber.open(io.BytesIO(raw)) as pdf:
        for page in pdf.pages:
            for table in page.extract_tables() or []:
                if not table or len(table) < 2:
                    continue
                header, *rows = table
                df = pd.DataFrame(rows, columns=header)
                entries.extend(_rows_to_entries(df))
    if entries:
        return entries

    with pdfplumber.open(io.BytesIO(raw)) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            for line in text.splitlines():
                line = line.strip()
                if not line:
                    continue
                match = re.search(r"(.+?)\s+(\(?-?\$?[\d,]+\.\d{2}\)?)\s*$", line)
                if not match:
                    continue
                name = match.group(1).strip()
                balance = _to_float(match.group(2))
                if name and balance is not None and not name.lower().startswith("total"):
                    entries.append(TBEntry(account_name=name, balance=balance))
    return entries


def parse_pasted_text(text: str) -> list[TBEntry]:
    text = text.strip()
    if not text:
        return []

    sniff_sample = text[:2048]
    delimiter = "\t" if "\t" in sniff_sample else ("," if "," in sniff_sample else None)
    if delimiter:
        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
        rows = list(reader)
        if rows:
            header_like = rows[0]
            looks_like_header = not any(_to_float(cell) is not None for cell in header_like)
            if looks_like_header:
                df = pd.DataFrame(rows[1:], columns=header_like)
            else:
                df = pd.DataFrame(rows)
            entries = _rows_to_entries(df)
            if entries:
                return entries

    entries: list[TBEntry] = []
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        match = re.search(r"(.+?)[\s,]+(\(?-?\$?[\d,]+\.?\d*\)?)\s*$", line)
        if not match:
            continue
        name = match.group(1).strip().strip(",")
        balance = _to_float(match.group(2))
        if name and balance is not None and name.lower() not in ("total", "totals"):
            entries.append(TBEntry(account_name=name, balance=balance))
    return entries


# ---------------------------------------------------------------------------
# Format auto-detection entry points
# ---------------------------------------------------------------------------


def inspect_upload(filename: str, raw: bytes) -> dict:
    """Inspect a file before full parsing: report detected format and, for audit
    workpapers, the available balance columns to choose from."""
    lower = filename.lower()
    if lower.endswith((".xlsx", ".xls")):
        audit_info = inspect_audit_workpaper(raw)
        if audit_info["is_audit_workpaper"]:
            return {"format": "audit_workpaper", **audit_info}
        if is_client_debit_credit_format(raw):
            return {"format": "client_debit_credit"}
        return {"format": "generic"}
    return {"format": "generic"}


def parse_upload(filename: str, raw: bytes, balance_column: str | None = None) -> list[TBEntry]:
    lower = filename.lower()
    if lower.endswith((".xlsx", ".xls")):
        audit_info = inspect_audit_workpaper(raw)
        if audit_info["is_audit_workpaper"]:
            return parse_audit_workpaper(raw, balance_column=balance_column)
        if is_client_debit_credit_format(raw):
            return parse_client_debit_credit(raw)
        return parse_excel_generic(raw)
    if lower.endswith(".csv"):
        return parse_csv(raw)
    if lower.endswith(".pdf"):
        return parse_pdf(raw)
    raise ValueError(f"Unsupported file type: {filename}")
