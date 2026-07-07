"""Test suite: parsing, matching (x - y sign convention), export, and API.

Run from backend/:  pip install pytest && pytest -q
The AI pass is exercised only as a no-op here (tests must run without an
ANTHROPIC_API_KEY so results are deterministic).
"""
from __future__ import annotations

import base64
import csv
import io
import json
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from main import app
from matching import ComparisonRow, build_comparison
from parsing import (
    ColumnMapping,
    TBEntry,
    _to_float,
    detect_client_name,
    detect_period_label,
    inspect_upload,
    parse_client_debit_credit,
    parse_pasted_text,
    parse_upload,
    parse_with_mapping,
)

EXAMPLES = Path(__file__).resolve().parent.parent / "examples"
AUDIT_XLSX = EXAMPLES / "Harborview Consulting LLC Working Trial Balance 12-31-2025.xlsx"
CLIENT_CSV = EXAMPLES / "Harborview Consulting LLC TB Detail 12-31-2025.csv"


@pytest.fixture(autouse=True)
def no_api_key(monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)


@pytest.fixture
def client():
    return TestClient(app)


# ---------------------------------------------------------------------------
# Difference convention: x = audit (Per Report), y = client, difference = x - y
# ---------------------------------------------------------------------------


def test_difference_is_audit_minus_client():
    row = ComparisonRow(None, "Professional Fees", client_balance=49760.0, audit_balance=51750.0, method="exact")
    assert row.difference == 1990.0


def test_difference_audit_only_row():
    row = ComparisonRow(None, "Interest Income", client_balance=None, audit_balance=-1250.0, method="audit_only")
    assert row.difference == -1250.0


def test_difference_client_only_row():
    row = ComparisonRow(None, "Bank Service Charges", client_balance=740.0, audit_balance=None, method="client_only")
    assert row.difference == -740.0


def test_materiality_threshold_is_one_dollar():
    under = ComparisonRow(None, "A", client_balance=100.0, audit_balance=100.99, method="exact")
    at = ComparisonRow(None, "B", client_balance=100.0, audit_balance=101.0, method="exact")
    assert not under.is_material
    assert at.is_material


def test_report_totals_follow_x_minus_y():
    internal = [TBEntry("Cash", 100.0), TBEntry("Extra Client Acct", 40.0)]
    audited = [TBEntry("Cash", 110.0), TBEntry("Extra Audit Acct", -25.0)]
    report = build_comparison(internal, audited, "Report")
    assert report.total_client == 140.0
    assert report.total_audit == 85.0
    # net difference = sum of per-row (audit - client) = total_audit - total_client
    assert report.net_difference == pytest.approx(report.total_audit - report.total_client)
    assert report.net_difference == -55.0


def test_ai_pass_failure_degrades_to_fuzzy_only(monkeypatch):
    # With an API key set but the API failing (timeout, bad key, network),
    # the reconcile must still succeed using the exact/fuzzy matches.
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test-key-not-real")

    def boom(*args, **kwargs):
        raise RuntimeError("simulated API failure")

    monkeypatch.setattr("matching._ai_pass_single", boom)
    internal = [TBEntry("Cash", 100.0), TBEntry("Weirdly Named Acct", 40.0)]
    audited = [TBEntry("Cash", 100.0), TBEntry("Completely Different", 40.0)]
    report = build_comparison(internal, audited, "Report")
    assert report.accounts_compared == 1
    assert report.only_in_client_count == 1
    assert report.only_in_audit_count == 1


def test_inspect_unknown_format_survives_ai_failure(monkeypatch, client):
    # An unrecognized file triggers the AI column-mapping suggestion inside
    # /api/inspect. With a key set but the API unreachable, inspect must still
    # answer quickly with format "unknown" (manual mapping panel) rather than
    # hanging or erroring — this was the "unable to fetch" on user uploads.
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test-key-not-real")
    monkeypatch.setenv("ANTHROPIC_BASE_URL", "http://127.0.0.1:9")  # connection refused
    csv_bytes = b"Col A,Col B\nSomething,123.45\nOther thing,67.89\n"
    resp = client.post("/api/inspect", files={"file": ("mystery.csv", csv_bytes, "text/csv")})
    assert resp.status_code == 200
    body = resp.json()
    assert body["format"] == "unknown"
    assert body["suggested_mapping"] is None
    assert body["grid_preview"]


def test_zero_totals_never_negative_zero():
    internal = [TBEntry("Cash", 0.005), TBEntry("AP", -0.005)]
    audited = [TBEntry("Cash", 0.005), TBEntry("AP", -0.005)]
    report = build_comparison(internal, audited, "Report")
    assert str(report.total_client) == "0.0"
    assert str(report.net_difference) == "0.0"


# ---------------------------------------------------------------------------
# Matching
# ---------------------------------------------------------------------------


def test_exact_match_ignores_punctuation_and_stopwords():
    report = build_comparison(
        [TBEntry("Furniture & Equipment", 1.0)], [TBEntry("Furniture and Equipment", 1.0)], "R"
    )
    assert report.rows[0].method == "exact"
    assert report.accounts_compared == 1


def test_fuzzy_match_near_identical_names():
    report = build_comparison(
        [TBEntry("Insurance Expenses", 5.0)], [TBEntry("Insurance Expense", 5.0)], "R"
    )
    assert report.rows[0].method == "fuzzy"


def test_dissimilar_names_stay_unmatched_without_ai():
    report = build_comparison(
        [TBEntry("Owner's Equity", -10.0)], [TBEntry("Members' Equity", -10.0)], "R"
    )
    methods = sorted(r.method for r in report.rows)
    assert methods == ["audit_only", "client_only"]
    assert report.only_in_client_count == 1
    assert report.only_in_audit_count == 1


def test_each_account_matched_at_most_once():
    internal = [TBEntry("Rent Expense", 1.0), TBEntry("Rent Expenses", 2.0)]
    audited = [TBEntry("Rent Expense", 1.0)]
    report = build_comparison(internal, audited, "R")
    matched = [r for r in report.rows if r.method in ("exact", "fuzzy")]
    assert len(matched) == 1
    assert matched[0].client_balance == 1.0  # the exact one wins


def test_rows_follow_audit_workpaper_order():
    internal = [TBEntry("B", 2.0), TBEntry("A", 1.0), TBEntry("Client Only", 9.0)]
    audited = [TBEntry("A", 1.0), TBEntry("Audit Only", 3.0), TBEntry("B", 2.0)]
    report = build_comparison(internal, audited, "R")
    assert [r.account_name for r in report.rows] == ["A", "Audit Only", "B", "Client Only"]


# ---------------------------------------------------------------------------
# Parsing
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "raw,expected",
    [
        ("1,234.56", 1234.56),
        ("(1,234.56)", -1234.56),
        ("$5,000", 5000.0),
        ("-42", -42.0),
        ("", None),
        ("-", None),
        ("N/A", None),
        (7, 7.0),
        (None, None),
    ],
)
def test_to_float(raw, expected):
    assert _to_float(raw) == expected


def test_parse_client_debit_credit_nets_and_stops_at_total():
    rows = [
        ["Full name", "Debit", "Credit"],
        ["Cash", "100.00", ""],
        ["Accounts Payable", "", "40.00"],
        ["TOTAL", "100.00", "40.00"],
        ["Junk after total", "999", ""],
    ]
    buf = io.StringIO()
    csv.writer(buf).writerows(rows)
    entries = parse_client_debit_credit("x.csv", buf.getvalue().encode())
    assert [(e.account_name, e.balance) for e in entries] == [("Cash", 100.0), ("Accounts Payable", -40.0)]


def test_parse_audit_workpaper_uses_tag_column():
    entries = parse_upload(AUDIT_XLSX.name, AUDIT_XLSX.read_bytes(), balance_column="Report 12/31/2025")
    names = [e.account_name for e in entries]
    assert len(entries) == 18
    assert "Cash and Cash Equivalents" in names
    # structural rows excluded
    assert not any("Total" in n or n in ("Assets", "Expenses") for n in names)
    assert round(sum(e.balance for e in entries), 2) == 0.0


def test_parse_audit_workpaper_heuristic_without_tag_column():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Code", "Account", "Description", "Adjusted Balance"])
    ws.append(["A", None, "Assets", None])  # header row: no Account -> skipped
    ws.append(["1000", "1000", "Cash", 100])
    ws.append(["1999 Total", "1999", "Total Assets", 100])  # 'Total' in code -> skipped
    buf = io.BytesIO()
    wb.save(buf)
    entries = parse_upload("wtb.xlsx", buf.getvalue(), balance_column="Adjusted Balance")
    assert [(e.account_name, e.balance) for e in entries] == [("Cash", 100.0)]


def test_parse_audit_workpaper_unknown_balance_column_raises():
    with pytest.raises(ValueError, match="not found"):
        parse_upload(AUDIT_XLSX.name, AUDIT_XLSX.read_bytes(), balance_column="Nope")


@pytest.mark.parametrize(
    "text",
    [
        "Cash\t100.00\nAccounts Payable\t-40.00",
        "Cash,100.00\nAccounts Payable,-40.00",
        "Account,Balance\nCash,100.00\nAccounts Payable,-40.00",
        "Cash 100.00\nAccounts Payable -40.00",
    ],
)
def test_parse_pasted_text_formats(text):
    entries = parse_pasted_text(text)
    assert [(e.account_name, e.balance) for e in entries] == [("Cash", 100.0), ("Accounts Payable", -40.0)]


def test_parse_with_mapping_debit_credit_and_stop_text():
    rows = [["ignored"], ["1", "Cash", "100", ""], ["2", "AP", "", "40"], ["", "STOP", "9", ""]]
    buf = io.StringIO()
    csv.writer(buf).writerows(rows)
    mapping = ColumnMapping(header_row=1, name_col=2, code_col=1, debit_col=3, credit_col=4, stop_text="stop")
    entries = parse_with_mapping("m.csv", buf.getvalue().encode(), mapping)
    assert [(e.account_code, e.account_name, e.balance) for e in entries] == [
        ("1", "Cash", 100.0),
        ("2", "AP", -40.0),
    ]


def test_column_mapping_rejects_bad_input():
    with pytest.raises(ValueError):
        ColumnMapping.from_dict({"name_col": 1})  # missing header_row
    with pytest.raises(ValueError):
        ColumnMapping.from_dict({"header_row": 1, "name_col": "abc"})


def test_inspect_upload_detects_each_format():
    audit = inspect_upload(AUDIT_XLSX.name, AUDIT_XLSX.read_bytes())
    assert audit["format"] == "audit_workpaper"
    assert audit["default_column"] == "Report 12/31/2025"
    assert audit["detected_date"] == "December 31, 2025"
    assert audit["detected_client_name"] == "Harborview Consulting, LLC"

    client = inspect_upload(CLIENT_CSV.name, CLIENT_CSV.read_bytes())
    assert client["format"] == "client_debit_credit"

    generic = inspect_upload("tb.csv", b"Account,Balance\nCash,1\n")
    assert generic["format"] == "generic"

    unknown = inspect_upload("mystery.csv", b"a,b\n1,2\n")
    assert unknown["format"] == "unknown"
    assert unknown["suggested_mapping"] is None  # no API key in tests
    assert unknown["grid_preview"]

    assert inspect_upload("notes.txt", b"hello")["format"] == "unsupported"


def test_xls_rejected_with_clear_message():
    # openpyxl can't read the legacy binary .xls format; without this guard the
    # user got a misleading "file may be corrupted" error.
    inspected = inspect_upload("old book.xls", b"\xd0\xcf\x11\xe0fake")
    assert inspected["format"] == "unsupported"
    assert "save it as .xlsx" in inspected["error"]
    with pytest.raises(ValueError, match=r"save it as \.xlsx"):
        parse_upload("old book.xls", b"\xd0\xcf\x11\xe0fake")


def test_detect_client_name_cases():
    assert detect_client_name("Acme Manufacturing LLC Trial Balance 12-31-2025.xlsx") == "Acme Manufacturing, LLC"
    assert detect_client_name("Widgets, Inc. WTB 2025-12-31.xlsx") == "Widgets, Inc."
    assert detect_client_name("randomfile.xlsx") is None


def test_detect_period_label_prefers_header_hint():
    assert detect_period_label("f 01-01-2020.xlsx", ["Unadjusted", "Report 12/31/2025"]) == "December 31, 2025"
    assert detect_period_label("tb 2025-06-30.csv") == "June 30, 2025"
    assert detect_period_label("no date here.csv") is None


# ---------------------------------------------------------------------------
# API + workbook export
# ---------------------------------------------------------------------------


def _reconcile_example(client) -> dict:
    result = client.post(
        "/api/reconcile",
        files={
            "client_file": (CLIENT_CSV.name, CLIENT_CSV.read_bytes()),
            "audit_file": (AUDIT_XLSX.name, AUDIT_XLSX.read_bytes()),
        },
        data={"client_name": "Harborview", "period_label": "Dec 31 2025", "balance_column": "Report 12/31/2025"},
    ).json()
    assert "error" not in result, result
    return result


def test_reconcile_example_end_to_end(client):
    data = _reconcile_example(client)
    s = data["summary"]
    assert s["accounts_compared"] == 14
    assert s["material_count"] == 1
    assert s["only_in_client_count"] == 4  # AI-pair accounts unmatched without a key
    assert s["only_in_audit_count"] == 4
    assert s["total_client"] == 0.0 and s["total_audit"] == 0.0 and s["net_difference"] == 0.0

    by_name = {r["account_name"]: r for r in data["rows"]}
    pf = by_name["Professional Fees"]
    assert (pf["audit_balance"], pf["client_balance"], pf["difference"]) == (51750.0, 49760.0, 1990.0)
    assert pf["is_material"]


def test_workbook_matches_json(client):
    data = _reconcile_example(client)
    wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(data["workbook_base64"])))
    assert wb.sheetnames == ["Adjusting Journal Entry", "TB Comparison", "Summary"]

    comp = wb["TB Comparison"]
    headers = [c.value for c in comp[3]]
    assert headers[2:5] == ["Per Client Records", "Per Report 12/31/2025", "Difference"]

    # client in C, audit in D, difference stays a live x - y formula
    body = list(comp.iter_rows(min_row=4, max_row=comp.max_row - 1))
    assert len(body) == len(data["rows"])
    for excel_row, json_row in zip(body, data["rows"]):
        assert excel_row[1].value == json_row["account_name"]
        assert excel_row[2].value == json_row["client_balance"]
        assert excel_row[3].value == json_row["audit_balance"]
        assert excel_row[4].value == f"=+D{excel_row[4].row}-C{excel_row[4].row}"

    total = list(comp.iter_rows(min_row=comp.max_row, max_row=comp.max_row))[0]
    last_data = comp.max_row - 1
    assert total[1].value == "TOTAL"
    assert total[2].value == f"=SUM(C4:C{last_data})"
    assert total[3].value == f"=SUM(D4:D{last_data})"
    assert total[4].value == f"=SUM(E4:E{last_data})"
    assert "0.00" in total[2].number_format  # zero totals stay visible

    labels = [row[0].value for row in wb["Summary"].iter_rows(min_row=7, max_row=14)]
    assert labels[-3:] == [
        "Total per client records (all accounts, signed)",
        "Total per audit working TB (all accounts, signed)",
        "Net difference across all accounts (audit − client)",
    ]


def test_workbook_totals_visible_without_recalculation(client):
    """Preview-only viewers (Drive/Gmail previews, Quick Look, LibreOffice with
    recalculation off) never run formulas — they show only cached results. Every
    formula cell must therefore carry its computed value, or the Difference
    column and the TOTAL rows render blank."""
    data = _reconcile_example(client)
    wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(data["workbook_base64"])), data_only=True)

    comp = wb["TB Comparison"]
    for excel_row, json_row in zip(comp.iter_rows(min_row=4, max_row=comp.max_row - 1), data["rows"]):
        assert excel_row[4].value == json_row["difference"]
    total = list(comp.iter_rows(min_row=comp.max_row, max_row=comp.max_row))[0]
    s = data["summary"]
    assert total[2].value == s["total_client"]
    assert total[3].value == s["total_audit"]
    assert total[4].value == s["net_difference"]

    aje = wb["Adjusting Journal Entry"]
    assert aje.cell(row=aje.max_row, column=2).value == data["aje"]["total_debit"]
    assert aje.cell(row=aje.max_row, column=3).value == data["aje"]["total_credit"]
    assert aje.cell(row=aje.max_row, column=4).value == 0.0  # balance check: debits = credits


def test_aje_lines_book_differences_as_debits_and_credits(client):
    data = _reconcile_example(client)
    aje = data["aje"]
    by_name = {r["account_name"]: r for r in aje["rows"]}

    # matched account with a difference books under the client's account name
    pf = by_name["Professional Fees"]
    assert (pf["debit"], pf["credit"]) == (1990.0, None)  # x - y = +1,990 -> debit
    # audit-only income account: negative difference -> credit
    assert by_name["Interest Income"]["credit"] == 1250.0
    # client-only expense account: bring to zero -> credit, under the client name
    assert by_name["Bank Service Charges"]["credit"] == 740.0
    # deep client path kept as-is (the client posts this entry in their books)
    assert "Operating Expenses:Occupancy:Facility Depreciation" in by_name

    # tied accounts produce no AJE line
    assert "Cash and Cash Equivalents" not in by_name
    # the entry balances when the net difference is zero
    assert aje["total_debit"] == aje["total_credit"] == 676240.0
    assert len(aje["rows"]) == 9


def test_aje_sheet_layout_and_formulas(client):
    data = _reconcile_example(client)
    wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(data["workbook_base64"])))
    ws = wb["Adjusting Journal Entry"]

    assert ws["A1"].value == "Harborview"
    assert ws["A2"].value == "Adjusting Journal Entry — To adjust client records to audited balances"
    assert ws["A3"].value == "As of Dec 31 2025"
    assert [c.value for c in ws[5][:3]] == ["Account / Description", "Debit", "Credit"]

    n = len(data["aje"]["rows"])
    first, last, totals_row = 6, 5 + n, 6 + n
    for row, json_line in zip(ws.iter_rows(min_row=first, max_row=last), data["aje"]["rows"]):
        assert row[0].value == json_line["account_name"]
        assert row[1].value == json_line["debit"]
        assert row[2].value == json_line["credit"]
        assert not (row[1].value and row[2].value)  # a line is a debit or a credit, never both

    assert ws.cell(row=totals_row, column=1).value == "TOTALS"
    assert ws.cell(row=totals_row, column=2).value == f"=SUM(B{first}:B{last})"
    assert ws.cell(row=totals_row, column=3).value == f"=SUM(C{first}:C{last})"
    assert ws.cell(row=totals_row, column=4).value == f"=+B{totals_row}-C{totals_row}"


def test_aje_empty_when_everything_ties(client):
    data = client.post(
        "/api/reconcile",
        data={"client_text": "Cash\t100.00", "audit_text": "Cash\t100.00",
              "client_name": "T", "period_label": "P"},
    ).json()
    assert data["aje"]["rows"] == []
    assert data["aje"]["total_debit"] == data["aje"]["total_credit"] == 0.0
    ws = openpyxl.load_workbook(io.BytesIO(base64.b64decode(data["workbook_base64"])))["Adjusting Journal Entry"]
    assert "No adjusting entries required" in ws["A6"].value
    assert ws["B7"].value == 0 and ws["C7"].value == 0  # totals are plain zeros, no SUM over nothing


def test_reconcile_pasted_text_both_sides(client):
    data = client.post(
        "/api/reconcile",
        data={
            "client_text": "Cash\t90.00",
            "audit_text": "Cash\t100.00",
            "client_name": "T",
            "period_label": "P",
        },
    ).json()
    assert data["rows"][0]["difference"] == 10.0  # x - y = 100 - 90


def test_reconcile_with_mapping_override(client):
    csv_bytes = b"z1,z2,z3\n1,Cash,100\n"
    mapping = json.dumps({"header_row": 1, "name_col": 2, "balance_col": 3, "code_col": 1})
    data = client.post(
        "/api/reconcile",
        files={"client_file": ("odd.csv", csv_bytes)},
        data={"audit_text": "Cash\t100.00", "client_mapping": mapping, "client_name": "T", "period_label": "P"},
    ).json()
    assert data["summary"]["accounts_compared"] == 1
    assert data["rows"][0]["difference"] == 0.0


def test_reconcile_rejects_empty_and_invalid_input(client):
    assert "error" in client.post("/api/reconcile", data={"client_name": "T"}).json()
    bad_mapping = client.post(
        "/api/reconcile",
        data={"client_text": "Cash\t1", "audit_text": "Cash\t1", "client_mapping": "{not json"},
    ).json()
    assert "error" in bad_mapping


def test_inspect_endpoint_handles_garbage(client):
    empty = client.post("/api/inspect", files={"file": ("x.xlsx", b"")}).json()
    assert empty["format"] == "error"
    corrupt = client.post("/api/inspect", files={"file": ("x.xlsx", b"not an xlsx")}).json()
    assert corrupt["format"] == "error"
    assert "Excel" in corrupt["error"]


def test_example_data_endpoint(client):
    data = client.get("/api/example-data").json()
    assert set(data) == {"client", "audit"}
    for side, path in (("client", CLIENT_CSV), ("audit", AUDIT_XLSX)):
        assert data[side]["filename"] == path.name
        assert base64.b64decode(data[side]["base64"]) == path.read_bytes()


def test_api_cache_headers(client):
    assert client.get("/api/example-data").headers["cache-control"] == "no-store"
    assert client.get("/").headers["cache-control"] == "no-cache"
