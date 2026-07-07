"""Build the exported workbook for a trial balance reconciliation.

Three sheets, mirroring the audit deliverable sent to the client for approval:
1. "Adjusting Journal Entry" — debit/credit lines that adjust the client's
   records to the audited balances, with SUM totals and a balance check cell.
2. "TB Comparison" — per-account detail (Per Client Records / Per Report /
   Difference-as-formula), TOTAL row with live SUM formulas.
3. "Summary" — health-check counts, totals, conclusion, and notes.
"""
from __future__ import annotations

import io
import re
import textwrap
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

from matching import ComparisonReport, MATERIAL_THRESHOLD

DARK_BLUE = "FF1F4E78"
WHITE = "FFFFFFFF"
GREEN = "FFC6EFCE"
RED = "FFFFC7CE"
YELLOW = "FFFFFF00"
TOTALS_FILL = "FFD9E1F2"
NUMBER_FORMAT = r"#,##0;\(#,##0\);\-"
# Totals must stay visible even when a signed trial balance sums to exactly zero,
# so this shows 0.00 (with cents) instead of the dash the account rows use.
TOTAL_NUMBER_FORMAT = r"#,##0.00;\(#,##0.00\);0.00"
AJE_NUMBER_FORMAT = "#,##0"

_HEADER_FONT = Font(bold=True, size=11, color=WHITE)
_HEADER_FILL = PatternFill(fill_type="solid", fgColor=DARK_BLUE)
_TITLE_FONT = Font(bold=True, size=14)
_SUBTITLE_FONT = Font(bold=True, size=11)
_BOLD_FONT = Font(bold=True)

_THIN = Side(style="thin")
_DOUBLE = Side(style="double")
_BOX_BORDER = Border(top=_THIN, bottom=_THIN, left=_THIN, right=_THIN)
_TOTALS_BORDER = Border(top=_DOUBLE, bottom=_DOUBLE, left=_THIN, right=_THIN)

_AJE_TITLE_FONT = Font(bold=True, size=14, color=DARK_BLUE)
_AJE_SUBTITLE_FONT = Font(bold=True, size=11, color="FF404040")
_AJE_PERIOD_FONT = Font(size=10, color="FF595959")
_AJE_HEADER_FONT = Font(bold=True, size=10, color=WHITE)
_AJE_BODY_FONT = Font(size=10)
_AJE_TOTALS_FONT = Font(bold=True, size=10)


def _header_row(ws: Worksheet, row: int, headers: list[str]) -> None:
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL


def _build_aje_sheet(
    ws: Worksheet, report: ComparisonReport, client_name: str, period_label: str
) -> dict[str, float]:
    ws.column_dimensions["A"].width = 56
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 13
    ws.column_dimensions["D"].width = 13

    for row in (1, 2, 3):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    ws["A1"] = client_name
    ws["A1"].font = _AJE_TITLE_FONT
    ws["A2"] = "Adjusting Journal Entry — To adjust client records to audited balances"
    ws["A2"].font = _AJE_SUBTITLE_FONT
    ws["A3"] = f"As of {period_label}"
    ws["A3"].font = _AJE_PERIOD_FONT

    header_row = 5
    for col, (text, align) in enumerate(
        [("Account / Description", "left"), ("Debit", "right"), ("Credit", "right")], start=1
    ):
        cell = ws.cell(row=header_row, column=col, value=text)
        cell.font = _AJE_HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.border = _BOX_BORDER
        cell.alignment = Alignment(horizontal=align)

    row = header_row + 1
    if not report.aje_rows:
        cell = ws.cell(
            row=row, column=1,
            value="No adjusting entries required — client records tie to the audited balances.",
        )
        cell.font = _AJE_BODY_FONT
        for col in (1, 2, 3):
            ws.cell(row=row, column=col).border = _BOX_BORDER
        row += 1
    for line in report.aje_rows:
        ws.cell(row=row, column=1, value=line.account_name).font = _AJE_BODY_FONT
        for col, value in ((2, line.debit), (3, line.credit)):
            cell = ws.cell(row=row, column=col, value=value)
            cell.font = _AJE_BODY_FONT
            cell.number_format = AJE_NUMBER_FORMAT
        for col in (1, 2, 3):
            ws.cell(row=row, column=col).border = _BOX_BORDER
        row += 1

    first_data, last_data = header_row + 1, row - 1
    label_cell = ws.cell(row=row, column=1, value="TOTALS")
    label_cell.alignment = Alignment(horizontal="right")
    total_debit = round(sum(line.debit or 0.0 for line in report.aje_rows), 2) + 0.0
    total_credit = round(sum(line.credit or 0.0 for line in report.aje_rows), 2) + 0.0
    cached_values: dict[str, float] = {}
    if report.aje_rows:
        ws.cell(row=row, column=2, value=f"=SUM(B{first_data}:B{last_data})")
        ws.cell(row=row, column=3, value=f"=SUM(C{first_data}:C{last_data})")
        cached_values[f"B{row}"] = total_debit
        cached_values[f"C{row}"] = total_credit
    else:
        ws.cell(row=row, column=2, value=0)
        ws.cell(row=row, column=3, value=0)
    for col in (1, 2, 3):
        cell = ws.cell(row=row, column=col)
        cell.font = _AJE_TOTALS_FONT
        cell.fill = PatternFill(fill_type="solid", fgColor=TOTALS_FILL)
        cell.border = _TOTALS_BORDER
        if col > 1:
            cell.number_format = AJE_NUMBER_FORMAT
    # Balance check just outside the table: debits minus credits must be zero.
    check = ws.cell(row=row, column=4, value=f"=+B{row}-C{row}")
    check.number_format = AJE_NUMBER_FORMAT
    cached_values[f"D{row}"] = round(total_debit - total_credit, 2) + 0.0
    return cached_values


def _build_summary_sheet(ws: Worksheet, report: ComparisonReport, client_name: str, period_label: str) -> None:
    ws.column_dimensions["A"].width = 92
    ws.column_dimensions["B"].width = 18

    ws["A1"] = client_name
    ws["A1"].font = _TITLE_FONT
    ws["A2"] = "Trial Balance Comparison — Client Records vs. Audit Working Trial Balance"
    ws["A2"].font = _SUBTITLE_FONT
    ws["A3"] = period_label
    ws["A4"] = f'Compared column: Audit Working Trial Balance "{report.compared_column}"'

    _header_row(ws, 6, ["Result", "Count / Amount"])

    metric_rows = [
        ("Accounts compared (present in both trial balances)", report.accounts_compared, None, NUMBER_FORMAT),
        ("Accounts that tie (difference under $1)", report.accounts_tied, None, NUMBER_FORMAT),
        ("Accounts with a material difference (≥ $1)", report.material_count, report.material_count == 0, NUMBER_FORMAT),
        ("Accounts with a balance only in client records", report.only_in_client_count, report.only_in_client_count == 0, NUMBER_FORMAT),
        ("Accounts with a balance only on audit working TB", report.only_in_audit_count, report.only_in_audit_count == 0, NUMBER_FORMAT),
        ("Total per client records (all accounts, signed)", report.total_client, abs(report.total_client) < MATERIAL_THRESHOLD, TOTAL_NUMBER_FORMAT),
        ("Total per audit working TB (all accounts, signed)", report.total_audit, abs(report.total_audit) < MATERIAL_THRESHOLD, TOTAL_NUMBER_FORMAT),
        ("Net difference across all accounts (audit − client)", report.net_difference, abs(report.net_difference) < MATERIAL_THRESHOLD, TOTAL_NUMBER_FORMAT),
    ]

    row = 7
    for label, value, is_good, number_format in metric_rows:
        ws.cell(row=row, column=1, value=label)
        value_cell = ws.cell(row=row, column=2, value=value)
        value_cell.number_format = number_format
        if is_good is not None:
            value_cell.fill = PatternFill(fill_type="solid", fgColor=GREEN if is_good else RED)
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Conclusion").font = _SUBTITLE_FONT
    row += 1
    for line in textwrap.wrap(report.conclusion, width=110):
        ws.cell(row=row, column=1, value=line)
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Notes").font = _SUBTITLE_FONT
    row += 1
    for note in report.notes:
        ws.cell(row=row, column=1, value=note)
        row += 1


def _build_comparison_sheet(
    ws: Worksheet, report: ComparisonReport, client_name: str, period_label: str
) -> dict[str, float]:
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 58
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 22
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["F"].width = 40

    ws["A1"] = f"{client_name} — Trial Balance Comparison as of {period_label}"
    ws["A1"].font = _TITLE_FONT

    header_row = 3
    _header_row(
        ws,
        header_row,
        [
            "Account Code",
            "Account / Description",
            "Per Client Records",
            f"Per {report.compared_column}",
            "Difference",
            "Notes",
        ],
    )

    row = header_row + 1
    cached_values: dict[str, float] = {}
    for r in report.rows:
        ws.cell(row=row, column=1, value=r.account_code)
        ws.cell(row=row, column=2, value=r.account_name)
        c_cell = ws.cell(row=row, column=3, value=r.client_balance)
        a_cell = ws.cell(row=row, column=4, value=r.audit_balance)
        # Difference stays a live formula (x - y: Per Report minus Per Client).
        d_cell = ws.cell(row=row, column=5, value=f"=+D{row}-C{row}")
        cached_values[f"E{row}"] = r.difference
        ws.cell(row=row, column=6, value=r.note)
        for cell in (c_cell, a_cell, d_cell):
            cell.number_format = NUMBER_FORMAT
        row += 1

    first_data, last_data = header_row + 1, row - 1
    if last_data >= first_data:
        ws.conditional_formatting.add(
            f"E{first_data}:E{last_data}",
            FormulaRule(
                formula=[f"ABS(E{first_data})>={MATERIAL_THRESHOLD}"],
                fill=PatternFill(fill_type="solid", fgColor=YELLOW),
            ),
        )

    total_cell = ws.cell(row=row, column=2, value="TOTAL")
    total_cell.font = _BOLD_FONT
    column_totals = {"C": report.total_client, "D": report.total_audit, "E": report.net_difference}
    for col_letter, col in (("C", 3), ("D", 4), ("E", 5)):
        cell = ws.cell(row=row, column=col, value=f"=SUM({col_letter}{first_data}:{col_letter}{last_data})")
        cell.number_format = TOTAL_NUMBER_FORMAT
        cell.font = _BOLD_FONT
        cached_values[f"{col_letter}{row}"] = column_totals[col_letter]
    return cached_values


_SHEET_ML_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def _inject_cached_formula_values(xlsx_bytes: bytes, sheet_values: dict[int, dict[str, float]]) -> bytes:
    """Store each formula cell's computed result in the sheet XML.

    openpyxl writes formulas without cached values, so viewers that don't run a
    calculation engine (Google Drive/Gmail previews, macOS Quick Look, LibreOffice
    with recalculation off) render the Difference column and every TOTAL row as
    blank cells. Embedding the precomputed result (the <v> element next to <f>)
    makes totals visible everywhere; Excel still recalculates the live formulas
    on open because openpyxl sets fullCalcOnLoad.

    sheet_values maps a zero-based sheet index (openpyxl saves worksheets in
    order as xl/worksheets/sheet1.xml, sheet2.xml, …) to {cell ref: value}.
    """
    source = zipfile.ZipFile(io.BytesIO(xlsx_bytes))
    buffer = io.BytesIO()
    ET.register_namespace("", _SHEET_ML_NS)
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            match = re.fullmatch(r"xl/worksheets/sheet(\d+)\.xml", item.filename)
            values = sheet_values.get(int(match.group(1)) - 1) if match else None
            if values:
                root = ET.fromstring(data)
                for cell in root.iter(f"{{{_SHEET_ML_NS}}}c"):
                    value = values.get(cell.get("r"))
                    if value is not None and cell.find(f"{{{_SHEET_ML_NS}}}f") is not None:
                        # openpyxl already writes an empty <v/> after <f> — fill it
                        # rather than appending a duplicate, which readers ignore.
                        v = cell.find(f"{{{_SHEET_ML_NS}}}v")
                        if v is None:
                            v = ET.SubElement(cell, f"{{{_SHEET_ML_NS}}}v")
                        # + 0.0 folds a rounded -0.0 into 0.0 so previews never show "-0".
                        v.text = repr(value + 0.0)
                data = ET.tostring(root, xml_declaration=True, encoding="UTF-8")
            target.writestr(item, data)
    return buffer.getvalue()


def build_workbook(report: ComparisonReport, client_name: str, period_label: str) -> bytes:
    wb = Workbook()
    aje_ws = wb.active
    aje_ws.title = "Adjusting Journal Entry"
    aje_values = _build_aje_sheet(aje_ws, report, client_name, period_label)

    comparison_ws = wb.create_sheet("TB Comparison")
    comparison_values = _build_comparison_sheet(comparison_ws, report, client_name, period_label)

    summary_ws = wb.create_sheet("Summary")
    _build_summary_sheet(summary_ws, report, client_name, period_label)

    buffer = io.BytesIO()
    wb.save(buffer)
    return _inject_cached_formula_values(buffer.getvalue(), {0: aje_values, 1: comparison_values})
